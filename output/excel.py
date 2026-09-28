"""Consulting-ready XLSX outputs built from Databank analytical results."""

from __future__ import annotations

from io import BytesIO
import re
from math import ceil

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ui.theme import GREY_LIGHT, PURPLE, WHITE


_INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_PRIMARY_FILL = PatternFill("solid", fgColor=PURPLE.removeprefix("#"))
_WHITE_FONT = Font(name="Arial", size=10, bold=True, color=WHITE.removeprefix("#"))
_BODY_FONT = Font(name="Arial", size=10, color="000000")
_TITLE_FONT = Font(name="Arial", size=16, bold=True, color=WHITE.removeprefix("#"))
_HEADER_BORDER = Border(bottom=Side(style="thin", color="FFFFFF"))

NUMBER_FORMAT = '#,##0;[Red](#,##0);-'
DECIMAL_FORMAT = '#,##0.0;[Red](#,##0.0);-'
PERCENTAGE_FORMAT = '0.0%;[Red](0.0%);-'
MULTIPLE_FORMAT = '0.0x;[Red](0.0x);-'
DKK_BILLION_TDK_FORMAT = '0.00,, "DKK mia.";[Red](0.00,, "DKK mia.");-'
DKK_BILLION_TDK_ONE_DECIMAL_FORMAT = '0.0,, "DKK mia.";[Red](0.0,, "DKK mia.");-'
_SECTION_LABELS = {
    "Earnings": "Indtjening",
    "Profitability": "Profitabilitet",
    "Efficiency": "Effektivitet",
    "Growth / balance sheet": "Vækst og balance",
}


def _blank_if_missing(value):
    return None if pd.isna(value) else value


def _new_sheet(workbook: Workbook, title: str, subtitle: str | None = None):
    worksheet = workbook.create_sheet(title)
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 26
    worksheet["A1"] = title
    worksheet["A1"].font = _TITLE_FONT
    worksheet["A1"].fill = _PRIMARY_FILL
    if subtitle:
        worksheet["A2"] = subtitle
        worksheet["A2"].font = Font(name="Arial", size=10, italic=True, color="5C5C5C")
    return worksheet


def _apply_title_band(worksheet, end_column: int) -> None:
    """Extend the compact Thursday title band only across the used sheet width."""
    for column in range(1, end_column + 1):
        cell = worksheet.cell(row=1, column=column)
        cell.fill = _PRIMARY_FILL
        cell.font = _TITLE_FONT


def _wrapped_row_height(value: object, width: float) -> float:
    if value is None:
        return 16
    return max(16, 15 * ceil(len(str(value)) / max(width * 1.15, 1)))


def _write_key_values(worksheet, values: list[tuple[str, object]], start_row: int = 4):
    for offset, (label, value) in enumerate(values):
        row = start_row + offset
        worksheet.cell(row=row, column=1, value=label).font = Font(name="Arial", size=10, bold=True)
        cell = worksheet.cell(row=row, column=2, value=_blank_if_missing(value))
        cell.font = _BODY_FONT
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            cell.alignment = Alignment(horizontal="right")
    return start_row + len(values) + 1


def _write_table(worksheet, frame: pd.DataFrame, *, start_row: int, number_formats=None, widths=None):
    number_formats = number_formats or {}
    widths = widths or {}
    for column_index, header in enumerate(frame.columns, start=1):
        cell = worksheet.cell(row=start_row, column=column_index, value=header)
        cell.font = _WHITE_FONT
        cell.fill = _PRIMARY_FILL
        cell.border = _HEADER_BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center")
        worksheet.column_dimensions[cell.column_letter].width = widths.get(header, min(max(len(str(header)) + 2, 12), 34))
    worksheet.row_dimensions[start_row].height = 20

    for row_index, row in enumerate(frame.itertuples(index=False, name=None), start=start_row + 1):
        wrapped_height = 16
        for column_index, (header, value) in enumerate(zip(frame.columns, row), start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=_blank_if_missing(value))
            cell.font = _BODY_FONT
            cell.alignment = Alignment(
                horizontal="right" if header in number_formats else "left",
                vertical="top",
                wrap_text=header in {"Beskrivelse", "Definition", "Eksklusionsårsag"},
            )
            if header in number_formats:
                cell.number_format = number_formats[header]
            if header in {"Beskrivelse", "Definition", "Eksklusionsårsag"}:
                wrapped_height = max(wrapped_height, _wrapped_row_height(value, worksheet.column_dimensions[cell.column_letter].width))
        worksheet.row_dimensions[row_index].height = wrapped_height
    worksheet.freeze_panes = f"A{start_row + 1}"


def _workbook_bytes(build_sheets) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    build_sheets(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _apply_metric_number_formats(worksheet, frame: pd.DataFrame, *, header_row: int, value_columns: list[str]):
    """Format metric values from their existing unit metadata, without conversion."""
    header_indexes = {header: index + 1 for index, header in enumerate(frame.columns)}
    for offset, (_, row) in enumerate(frame.iterrows(), start=header_row + 1):
        number_format = {
            "percentage": PERCENTAGE_FORMAT,
            "multiple": MULTIPLE_FORMAT,
            "DKK mia.": DKK_BILLION_TDK_FORMAT,
        }.get(row["Enhed"], DECIMAL_FORMAT)
        for header in value_columns:
            worksheet.cell(row=offset, column=header_indexes[header]).number_format = number_format


def insurance_market_structure_workbook(summary: pd.DataFrame, source_table: pd.DataFrame) -> bytes:
    """Export the supplied market-structure engine results without recalculation."""
    summary_table = summary.loc[:, ["year", "entity_count", "market_size", "cr1", "cr3", "cr5", "hhi", "known_data_break"]].rename(
        columns={
            "year": "År", "entity_count": "Enheder med positive bruttopræmier", "market_size": "Markedsstørrelse (t.DKK)",
            "cr1": "CR1", "cr3": "CR3", "cr5": "CR5", "hhi": "HHI", "known_data_break": "Kendt databrud",
        }
    )
    shares_table = source_table.loc[:, ["year", "display_name", "entity_id", "market_value", "market_share", "rank", "included_flag", "exclusion_reason"]].rename(
        columns={
            "year": "År", "display_name": "Selskab", "entity_id": "Enheds-ID", "market_value": "Markedsværdi (t.DKK)",
            "market_share": "Markedsandel", "rank": "Rang", "included_flag": "Indgår", "exclusion_reason": "Eksklusionsårsag",
        }
    )
    period = f"Valgt periode: {int(summary['year'].min())}-{int(summary['year'].max())}" if not summary.empty else "Valgt periode"
    methodology = pd.DataFrame(
        [
            ("Markedsgrundlag", "Res_BP_BeY / Bruttopræmier"), ("Enhed", "t.DKK"),
            ("Population", "Kun juridiske enheder med observerede positive bruttopræmier."),
            ("Koncentration", "CR1, CR3 og CR5 er kumulerede markedsandele; HHI er på skalaen 0-10.000."),
            ("Databrud", "2025 er et kendt dækningsbrud og bør fortolkes med forsigtighed."),
            ("Metode", "Se docs/MARKET_STRUCTURE_METHOD.md."),
        ], columns=["Emne", "Beskrivelse"],
    )

    def build(workbook):
        overview = _new_sheet(workbook, "Overblik", "Insurance Market Structure")
        _write_key_values(overview, [("Periode", period.removeprefix("Valgt periode: "))])
        _write_table(overview, summary_table, start_row=6, number_formats={"År": NUMBER_FORMAT, "Enheder med positive bruttopræmier": NUMBER_FORMAT, "Markedsstørrelse (t.DKK)": NUMBER_FORMAT, "CR1": PERCENTAGE_FORMAT, "CR3": PERCENTAGE_FORMAT, "CR5": PERCENTAGE_FORMAT, "HHI": NUMBER_FORMAT}, widths={"Enheder med positive bruttopræmier": 34, "Markedsstørrelse (t.DKK)": 25, "Kendt databrud": 18})
        _apply_title_band(overview, len(summary_table.columns))
        shares = _new_sheet(workbook, "Markedsandele", period)
        _write_table(shares, shares_table, start_row=4, number_formats={"År": NUMBER_FORMAT, "Markedsværdi (t.DKK)": NUMBER_FORMAT, "Markedsandel": PERCENTAGE_FORMAT, "Rang": NUMBER_FORMAT}, widths={"Selskab": 34, "Enheds-ID": 20, "Markedsværdi (t.DKK)": 24, "Eksklusionsårsag": 30})
        _apply_title_band(shares, len(shares_table.columns))
        method = _new_sheet(workbook, "Metode", "Kilde og afgrænsning")
        _write_table(method, methodology, start_row=4, widths={"Emne": 22, "Beskrivelse": 70})
        _apply_title_band(method, len(methodology.columns))

    return _workbook_bytes(build)


def bank_analyst_workbook(comparison: pd.DataFrame, history: pd.DataFrame, overview: dict, fiscal_year: int, benchmark_definition: str) -> bytes:
    """Export the supplied Bank Analyst View results without recalculation."""
    overview_values = [
        ("Bank", overview["display_name"]), ("Valgt finansår", fiscal_year), ("Benchmark", benchmark_definition),
        ("Aktiver i alt (DKK mia.)", overview["total_assets"]), ("Kernemetrikker med data", overview["metrics_with_data"]),
        ("Regnr", overview["regnr"]), ("Enheds-ID", overview["entity_id"]),
    ]
    metrics = comparison.loc[:, ["section", "display_name", "current_value", "previous_value", "yoy_change", "peer_median", "unit", "display_format"]].assign(
        section=lambda frame: frame["section"].map(_SECTION_LABELS),
        unit=lambda frame: frame["display_format"].map({"dkk_billion_tdk": "DKK mia."}).fillna(frame["unit"]),
    ).rename(
        columns={"section": "Sektion", "display_name": "Metrik", "current_value": "Aktuelt år", "previous_value": "Foregående år", "yoy_change": "YoY", "peer_median": "Benchmarkmedian", "unit": "Enhed"}
    )
    metrics = metrics[["Sektion", "Metrik", "Aktuelt år", "Foregående år", "YoY", "Benchmarkmedian", "Enhed"]]
    metric_names = comparison.set_index("metric_id")["display_name"]
    metric_units = comparison.set_index("metric_id")["unit"]
    metric_display_formats = comparison.set_index("metric_id")["display_format"]
    history_table = history.loc[:, ["metric_id", "fiscal_year", "value"]].assign(
        Metrik=lambda frame: frame["metric_id"].map(metric_names),
        Enhed=lambda frame: frame["metric_id"].map(metric_display_formats).map({"dkk_billion_tdk": "DKK mia."}).fillna(frame["metric_id"].map(metric_units)),
    )[["Metrik", "fiscal_year", "value", "Enhed"]].rename(columns={"fiscal_year": "Finansår", "value": "Værdi"})
    technical = comparison.loc[:, ["metric_id", "display_name", "source_type", "calculation_type", "validation_status", "display_format", "definition"]].rename(
        columns={"metric_id": "Metric ID", "display_name": "Metrik", "source_type": "Kildetype", "calculation_type": "Beregningstype", "validation_status": "Valideringsstatus", "display_format": "Visningsformat", "definition": "Definition"}
    )
    methodology = pd.DataFrame(
        [
            ("YoY", "Aktuel værdi minus samme metriks observerede værdi i umiddelbart foregående finansår."),
            ("Benchmark", f"{benchmark_definition}. Målbanken indgår, når den har data."),
            ("Identitet", "Kanonisk entity_id er baseret på regnr; displaynavn er kun en visningsattribut."),
            ("Manglende data", "Manglende observationer forbliver manglende og erstattes ikke med nul."),
            ("Proveniens", "Rapporterede og beregnede metrikker fremgår særskilt af Teknisk."),
            ("Udskudt", "Indlån, cost-income, indtjeningsvækst, kapitalprocenter og risikomål."),
            ("Metode", "Se docs/BANK_ANALYST_METHOD.md."),
        ], columns=["Emne", "Beskrivelse"],
    )

    def build(workbook):
        overview_sheet = _new_sheet(workbook, "Overblik", "Bank Analyst View")
        _write_key_values(overview_sheet, overview_values)
        overview_sheet["B7"].number_format = DKK_BILLION_TDK_ONE_DECIMAL_FORMAT
        overview_sheet["B8"].number_format = NUMBER_FORMAT
        overview_sheet.column_dimensions["A"].width = 28
        overview_sheet.column_dimensions["B"].width = 52
        _apply_title_band(overview_sheet, 2)
        metrics_sheet = _new_sheet(workbook, "Metrikker", f"{overview['display_name']} · FY {fiscal_year}")
        _write_table(metrics_sheet, metrics, start_row=4, number_formats={"Aktuelt år": DECIMAL_FORMAT, "Foregående år": DECIMAL_FORMAT, "YoY": DECIMAL_FORMAT, "Benchmarkmedian": DECIMAL_FORMAT}, widths={"Sektion": 23, "Metrik": 36, "Aktuelt år": 18, "Foregående år": 18, "Benchmarkmedian": 20, "Enhed": 14})
        _apply_metric_number_formats(metrics_sheet, metrics, header_row=4, value_columns=["Aktuelt år", "Foregående år", "YoY", "Benchmarkmedian"])
        _apply_title_band(metrics_sheet, len(metrics.columns))
        history_sheet = _new_sheet(workbook, "Historik", f"{overview['display_name']} · op til fem tilgængelige finansår")
        _write_table(history_sheet, history_table, start_row=4, number_formats={"Finansår": NUMBER_FORMAT, "Værdi": DECIMAL_FORMAT}, widths={"Metrik": 36, "Finansår": 14, "Værdi": 20, "Enhed": 14})
        _apply_metric_number_formats(history_sheet, history_table, header_row=4, value_columns=["Værdi"])
        _apply_title_band(history_sheet, len(history_table.columns))
        method_sheet = _new_sheet(workbook, "Metode", "Kilde og afgrænsning")
        _write_table(method_sheet, methodology, start_row=4, widths={"Emne": 22, "Beskrivelse": 70})
        _apply_title_band(method_sheet, len(methodology.columns))
        technical_sheet = _new_sheet(workbook, "Teknisk", "Teknisk metadata til reproducerbarhed")
        _write_table(technical_sheet, technical, start_row=4, widths={"Metric ID": 32, "Metrik": 36, "Valideringsstatus": 32, "Definition": 72})
        _apply_title_band(technical_sheet, len(technical.columns))

    return _workbook_bytes(build)


def insurance_market_structure_filename(start_year: int, end_year: int) -> str:
    return f"Databank_Insurance_Market_Structure_{start_year}-{end_year}.xlsx"


def bank_analyst_filename(display_name: str, fiscal_year: int) -> str:
    """Return a deterministic Windows-safe bank export filename."""
    safe_name = _INVALID_FILENAME_CHARS.sub("_", str(display_name))
    safe_name = re.sub(r"\s+", "_", safe_name).strip("._ ") or "Bank"
    return f"Databank_Bank_Analyst_{safe_name}_{fiscal_year}.xlsx"
