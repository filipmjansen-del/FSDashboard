from io import BytesIO
from pathlib import Path
import unittest

import pandas as pd
from openpyxl import load_workbook

from analytics.bank_analyst import bank_entities_for_year, build_bank_analyst_view
from analytics.insurance_market_structure import build_market_structure_table, summarize_market_structure
from data.access import load_insurance_market_structure_fp
from data_loader import load_raw_data
from output.excel import (
    bank_analyst_filename,
    bank_analyst_workbook,
    insurance_market_structure_filename,
    insurance_market_structure_workbook,
)


class OutputExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw = load_raw_data(Path(__file__).parents[1] / "financial_services_long.xlsx")

    def test_insurance_export_preserves_fp_engine_tables(self):
        source_table = build_market_structure_table(load_insurance_market_structure_fp())
        summary = summarize_market_structure(source_table)
        source_table = source_table.loc[source_table["quarter"].eq(2)].copy()
        workbook = load_workbook(BytesIO(insurance_market_structure_workbook(summary.loc[summary["quarter"].eq(2)], source_table)), data_only=True)

        self.assertEqual(workbook.sheetnames, ["Overblik", "Markedsandele", "Metode"])
        summary_sheet = workbook["Overblik"]
        self.assertEqual([cell.value for cell in summary_sheet[6]], [
            "År", "Kvartal", "Markedsaktører", "Bruttopræmieindtægter (t.DKK)", "CR1", "CR3", "CR5", "HHI",
        ])
        self.assertEqual(summary_sheet["D7"].value, summary.loc[summary["quarter"].eq(2)].iloc[0]["market_size"])
        self.assertEqual(summary_sheet["A6"].font.name, "Arial")
        self.assertEqual(summary_sheet["A6"].fill.fgColor.rgb[-6:], "412B48")
        self.assertEqual(summary_sheet["H1"].fill.fgColor.rgb[-6:], "412B48")
        self.assertFalse(summary_sheet.sheet_view.showGridLines)
        self.assertEqual(summary_sheet.freeze_panes, "A7")
        self.assertEqual(summary_sheet["A7"].number_format, "0")
        shares_sheet = workbook["Markedsandele"]
        self.assertEqual([cell.value for cell in shares_sheet[4]], [
            "År", "Kvartal", "Selskab", "Bruttopræmieindtægter (t.DKK)", "Rapporteret markedsandel", "Rang",
        ])
        self.assertEqual(shares_sheet["D5"].number_format, '#,##0;[Red](#,##0);-')

    def test_bank_export_preserves_exact_analyst_comparison_and_history(self):
        entities = bank_entities_for_year(self.raw, 2025)
        comparison, history, overview = build_bank_analyst_view(
            self.raw, "bank:3000", 2025, entities["entity_id"].tolist()
        )
        workbook = load_workbook(
            BytesIO(bank_analyst_workbook(comparison, history, overview, 2025, "Alle banker med en observation i det valgte år")),
            data_only=True,
        )

        self.assertEqual(workbook.sheetnames, ["Overblik", "Metrikker", "Historik", "Metode", "Teknisk"])
        comparison_sheet = workbook["Metrikker"]
        headers = [cell.value for cell in comparison_sheet[4]]
        self.assertEqual(headers, ["Sektion", "Metrik", "Aktuelt år", "Foregående år", "YoY", "Benchmarkmedian", "Enhed"])
        self.assertNotIn("Metric ID", headers)
        profit_row = comparison.loc[comparison["metric_id"].eq("bank.profit_before_tax")].iloc[0]
        exported_row = next(
            row for row in comparison_sheet.iter_rows(min_row=5, values_only=True)
            if row[headers.index("Metrik")] == profit_row["display_name"]
        )
        self.assertEqual(exported_row[headers.index("Aktuelt år")], profit_row["current_value"])
        self.assertEqual(exported_row[headers.index("Benchmarkmedian")], profit_row["peer_median"])
        self.assertEqual(comparison_sheet["C5"].number_format, '0.00,, "DKK mia.";[Red](0.00,, "DKK mia.");-')
        financial_row = next(
            row for row in comparison_sheet.iter_rows(min_row=5)
            if row[headers.index("Metrik")].value == "Netto renteindtægter"
        )
        self.assertEqual(financial_row[headers.index("Enhed")].value, "DKK mia.")
        self.assertEqual(financial_row[headers.index("Aktuelt år")].number_format, '0.00,, "DKK mia.";[Red](0.00,, "DKK mia.");-')
        unit_values = [row[headers.index("Enhed")].value for row in comparison_sheet.iter_rows(min_row=5)]
        self.assertIn("%", unit_values)
        self.assertIn("x", unit_values)
        self.assertNotIn("percentage", unit_values)
        self.assertNotIn("multiple", unit_values)
        self.assertEqual(comparison_sheet["G1"].fill.fgColor.rgb[-6:], "412B48")
        self.assertEqual(workbook["Overblik"]["B4"].alignment.horizontal, "left")
        self.assertEqual(workbook["Overblik"]["B5"].alignment.horizontal, "left")
        self.assertEqual(comparison_sheet["A4"].font.name, "Arial")
        self.assertFalse(comparison_sheet.sheet_view.showGridLines)
        self.assertEqual(comparison_sheet.freeze_panes, "A5")
        self.assertIn("Metric ID", [cell.value for cell in workbook["Teknisk"][4]])
        history_sheet = workbook["Historik"]
        self.assertEqual(history_sheet.max_row - 4, len(history))
        self.assertEqual(history_sheet["B5"].number_format, "0")
        self.assertTrue(all(workbook["Metode"].row_dimensions[row].height == 16 for row in range(5, workbook["Metode"].max_row + 1)))
        self.assertTrue(all(workbook["Teknisk"].row_dimensions[row].height == 16 for row in range(5, workbook["Teknisk"].max_row + 1)))
        self.assertGreater(workbook["Metode"].column_dimensions["B"].width, 100)
        self.assertGreater(workbook["Teknisk"].column_dimensions["G"].width, 100)

    def test_filenames_are_deterministic_and_safe(self):
        self.assertEqual(insurance_market_structure_filename(2), "Databank_Insurance_Market_Structure_Q2.xlsx")
        self.assertEqual(bank_analyst_filename('A/B: Bank', 2025), "Databank_Bank_Analyst_A_B__Bank_2025.xlsx")


if __name__ == "__main__":
    unittest.main()
